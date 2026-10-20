#!/usr/bin/env python3
"""Render selected pages of a PDF/PPTX to PNG images and extract their text.

Usage: python render_slides.py FILE --pages "all" | "80, 89-100" --out OUTDIR [--zoom 1.5]
Prints JSON: {"page_count": N, "pages": [{"page": 89, "image": "...png", "text": "..."}]}
Needs: pip install pymupdf   (PPTX also needs LibreOffice `soffice`, or python-pptx for text only)
"""
import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from common import ensure_dir, parse_pages  # noqa: E402


def to_pdf(path: Path) -> Path | None:
    if path.suffix.lower() == ".pdf":
        return path
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        out = Path(tempfile.mkdtemp())
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=300)
        return out / (path.stem + ".pdf")
    try:  # PowerPoint on Windows
        import comtypes.client
        out = Path(tempfile.mkdtemp()) / (path.stem + ".pdf")
        app = comtypes.client.CreateObject("PowerPoint.Application")
        pres = app.Presentations.Open(str(path.resolve()), WithWindow=False)
        pres.SaveAs(str(out), 32)
        pres.Close()
        app.Quit()
        return out
    except Exception:
        return None


def text_only_pptx(path: Path, spec: str) -> dict:
    from pptx import Presentation
    slides = Presentation(str(path)).slides
    pages = parse_pages(spec, len(slides))
    res = []
    for p in pages:
        parts = [sh.text_frame.text for sh in slides[p - 1].shapes if sh.has_text_frame]
        res.append({"page": p, "image": None, "text": "\n".join(parts)})
    return {"page_count": len(slides), "pages": res,
            "warning": "No LibreOffice/PowerPoint found: text only, no images. Diagrams and formulas may be missing."}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--pages", default="all")
    ap.add_argument("--out", required=True)
    ap.add_argument("--zoom", type=float, default=1.5)
    a = ap.parse_args()

    src = Path(a.file)
    pdf = to_pdf(src)
    if pdf is None:
        print(json.dumps(text_only_pptx(src, a.pages), indent=2))
        return

    try:
        import pymupdf as fitz
    except ImportError:  # older PyMuPDF
        import fitz
    doc = fitz.open(pdf)
    pages = parse_pages(a.pages, len(doc))
    out = ensure_dir(a.out)
    res = []
    for p in pages:
        page = doc[p - 1]
        img = out / f"slide_{p:03d}.png"
        page.get_pixmap(matrix=fitz.Matrix(a.zoom, a.zoom)).save(str(img))
        res.append({"page": p, "image": str(img), "text": page.get_text().strip()})
    print(json.dumps({"page_count": len(doc), "pages": res}, indent=2))


if __name__ == "__main__":
    main()
